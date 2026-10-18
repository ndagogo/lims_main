"""Enterprise reporting dashboard, catalogue, and export audit."""
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import Client, TestCase

from billing.models import Invoice, InvoiceItem, Payment, Refund
from core.models import AuditEvent, Centre, UserCentreAccess
from laboratory.models import ServiceGroup, Tariff, Test, TestPrice
from patients.models import OrganisationClient, Patient, Visit
from referrals.models import CommissionLedgerEntry, ReferralAttribution, ReferralDoctor
from reporting.metrics import financial_kpis
from reporting.tasks import _write_export_file

User = get_user_model()


class ReportingEnterpriseTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='acct1', email='acct1@example.test', password='Strong-test-password-2026!',
        )
        group = Group.objects.get_or_create(name='Accounts')[0]
        self.user.groups.add(group)
        self.centre = Centre.objects.create(code='NWS', name='NOWAS')
        UserCentreAccess.objects.create(
            user=self.user, centre=self.centre, can_view=True, can_edit=True, can_register=True,
        )
        self.patient = Patient.objects.create(
            first_name='Ada', last_name='Okeke', phone_number='8011112222', sex='Female',
        )
        self.visit = Visit.objects.create(
            patient=self.patient, centre=self.centre, registered_by=self.user, referral_source='DOCTOR',
        )
        group_lab = ServiceGroup.objects.create(name='Haematology', code='1')
        self.test = Test.objects.create(code='CBC', name='Complete Blood Count', service_group=group_lab)
        self.tariff = Tariff.objects.create(name='Standard')
        TestPrice.objects.create(tariff=self.tariff, test=self.test, price=Decimal('1000.00'))
        self.invoice = Invoice.objects.create(
            visit=self.visit, discount_amount=Decimal('100.00'), status='UNPAID',
        )
        InvoiceItem.objects.create(
            invoice=self.invoice, test=self.test, tariff=self.tariff,
            description='CBC', quantity=1, unit_price=Decimal('1000.00'),
        )
        self.invoice.refresh_from_db()
        self.payment = Payment.objects.create(
            invoice=self.invoice, amount=Decimal('500.00'), method='CASH',
            status='CONFIRMED', received_by=self.user,
        )
        self.invoice.refresh_from_db()
        self.client = Client()
        self.client.force_login(self.user)

    def test_financial_kpis_are_precise(self):
        kpis = financial_kpis(self.user, period='annual')['kpis']
        self.assertEqual(kpis['gross_charges'], Decimal('1000.00'))
        self.assertEqual(kpis['discounts'], Decimal('100.00'))
        self.assertEqual(kpis['invoice_value'], Decimal('900.00'))
        self.assertEqual(kpis['collections'], Decimal('500.00'))
        self.assertEqual(kpis['refunds'], Decimal('0.00'))
        self.assertEqual(kpis['net_receipts'], Decimal('500.00'))
        self.assertEqual(kpis['outstanding_receivables'], Decimal('400.00'))

    def test_management_dashboard_shows_four_panels(self):
        page = self.client.get('/reporting/management/?period=annual')
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'GROSS CHARGES')
        self.assertContains(page, 'NET RECEIPTS')
        self.assertContains(page, 'Clinical operations')
        self.assertContains(page, 'Centres')
        self.assertContains(page, 'Referral network')
        self.assertContains(page, '1000.00')
        self.assertContains(page, '500.00')

    def test_catalogue_revenue_by_centre_and_export_audited(self):
        page = self.client.get('/reporting/catalogue/?report=revenue_centre&period=annual')
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'Revenue by centre')
        self.assertContains(page, 'NWS')
        exported = self.client.get(
            '/reporting/export/?kind=catalogue&report=revenue_centre&format=csv&period=annual',
        )
        self.assertEqual(exported.status_code, 200)
        self.assertIn('text/csv', exported['Content-Type'])
        audit = AuditEvent.objects.filter(action='REPORT_EXPORTED').latest('created_at')
        self.assertEqual(audit.detail['report_type'], 'revenue_centre')
        self.assertFalse(audit.detail['include_patient_ids'])

    def test_hmo_statement_and_commission_catalogue(self):
        partner = OrganisationClient.objects.create(
            centre=self.centre, kind='HMO', name='AXA Mansard',
        )
        self.visit.client = partner
        self.visit.save(update_fields=['client'])
        doctor = ReferralDoctor.objects.create(
            centre=self.centre, doctor_code='NWS-DR9', first_name='Ife', last_name='Balogun',
            commission_type='PERCENT', commission_value=Decimal('5'),
        )
        ReferralAttribution.objects.create(visit=self.visit, doctor=doctor, recorded_by=self.user)
        CommissionLedgerEntry.objects.create(
            doctor=doctor, invoice=self.invoice, entry_type='ACCRUAL',
            amount=Decimal('45.00'), description='Test accrual', created_by=self.user,
        )
        hmo = self.client.get('/reporting/catalogue/?report=hmo_statements&period=annual')
        self.assertEqual(hmo.status_code, 200)
        self.assertContains(hmo, 'AXA Mansard')
        commission = self.client.get('/reporting/catalogue/?report=commission&period=annual')
        self.assertEqual(commission.status_code, 200)
        self.assertContains(commission, 'NWS-DR9')

    def test_management_sections_export_transaction_detail(self):
        from io import BytesIO

        from openpyxl import load_workbook

        marketer = User.objects.create_user(
            username='marketer1', first_name='Ada', last_name='Marketing',
            email='marketer1@example.test', password='Strong-test-password-2026!',
        )
        doctor = ReferralDoctor.objects.create(
            centre=self.centre, doctor_code='NWS-DR77', title='Dr',
            first_name='Ngozi', last_name='Eze', commission_type='PERCENT',
            commission_value=Decimal('5'), onboarded_by=marketer,
        )
        ReferralAttribution.objects.create(visit=self.visit, doctor=doctor, recorded_by=self.user)
        CommissionLedgerEntry.objects.create(
            doctor=doctor, invoice=self.invoice, entry_type='ACCRUAL',
            amount=Decimal('45.00'), description='Commission on invoice', created_by=self.user,
        )
        query = (
            '/reporting/management/?period=annual'
            '&section=financial&section=collections&section=centres'
            '&section=referral&section=physicians&section=introducers&section=clinical'
        )
        page = self.client.get(query)
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'Detailed reports')
        self.assertContains(page, self.invoice.invoice_number)
        self.assertContains(page, 'Cash')
        self.assertContains(page, 'Haematology')
        self.assertContains(page, 'Dr Ngozi Eze')
        self.assertContains(page, self.visit.visit_id)
        self.assertNotContains(page, 'Okeke')
        self.assertContains(page, 'No rows for this period and centre.')

        exported = self.client.get(
            '/reporting/export/?kind=management&detail=1&format=xlsx&period=annual'
            '&section=financial&section=collections',
        )
        self.assertEqual(exported.status_code, 200)
        book = load_workbook(BytesIO(exported.content))
        self.assertEqual(book.sheetnames, ['Financial revenue', 'Collections by method'])
        methods = [cell.value for cell in book['Collections by method']['E']]
        self.assertIn('Cash', methods)
        audit = AuditEvent.objects.filter(action='REPORT_EXPORTED').latest('created_at')
        self.assertEqual(audit.detail['report_type'], 'management_financial_collections')
        self.assertFalse(audit.detail['include_patient_ids'])

        missing = self.client.get(
            '/reporting/export/?kind=management&detail=1&format=xlsx&period=annual',
        )
        self.assertEqual(missing.status_code, 302)

    def test_background_export_writes_file(self):
        result = _write_export_file('management_summary', self.user.pk, {'period': 'annual', 'format': 'csv'})
        self.assertTrue(result['ok'])
        self.assertTrue(result['path'].endswith('.csv'))
        self.assertTrue(AuditEvent.objects.filter(action='REPORT_EXPORT_COMPLETED').exists())

    def test_net_receipts_subtract_refunds(self):
        Refund.objects.create(
            invoice=self.invoice, payment=self.payment, amount=Decimal('50.00'),
            reason='Partial refund', status='COMPLETED', authorised_by=self.user,
        )
        kpis = financial_kpis(self.user, period='annual')['kpis']
        self.assertEqual(kpis['refunds'], Decimal('50.00'))
        self.assertEqual(kpis['net_receipts'], Decimal('450.00'))
